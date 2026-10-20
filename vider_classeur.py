#!/usr/bin/env python3
"""
vider_classeur.py — Transforme un classeur d'essai « rempli » en fichier brut,
réinjectable dans l'app de calcul.

Règles :
  • Feuille DIAMETRE (CR,CS) :
        - conservé : en-têtes + A (ECH) + D_1, D_2, D_3
        - vidé     : D_Moy(mm), S_(mm2), D_Moy Global(mm)
  • Feuilles COROUI / COSOU :
        - conservé : en-têtes + Elong.(mm) + Force(N)
        - vidé     : Def, Cont, Cont(max), Def rupt, E(mpa)
  • Feuille RECCAPITULATIF :
        - vidée (ou supprimée avec --supprimer-recap)
  • Graphiques : supprimés (l'exporteur les régénère)

Usage :
    python vider_classeur.py RT4.xlsx
    python vider_classeur.py RT4.xlsx -o RT4_BRUT.xlsx
    python vider_classeur.py *.xlsx -d bruts/
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell


# ---------------------------------------------------------------------------
# Nettoyage / détection
# ---------------------------------------------------------------------------

def _norm(valeur) -> str:
    """Majuscules + suppression de tout ce qui n'est pas alphanumérique."""
    return re.sub(r'[^A-Z0-9]', '', str(valeur or '').upper())


def _classer(nom_feuille: str) -> str:
    n = _norm(nom_feuille)
    if n.startswith(('DIAMETRE', 'DIAMETER')):
        return 'diametre'
    if n in ('COROUI', 'COSOU'):
        return 'courbe'
    if n.startswith(('RECAPITULATIF', 'RECCAPITULATIF', 'RECAP')):
        return 'recap'
    return 'inconnu'


def _effacer(ws, row: int, col: int) -> None:
    """Efface une cellule (y compris si fusionnée)."""
    cell = ws.cell(row, col)
    if isinstance(cell, MergedCell):
        for rng in ws.merged_cells.ranges:
            if rng.min_row <= row <= rng.max_row and rng.min_col <= col <= rng.max_col:
                ws.cell(rng.min_row, rng.min_col).value = None
                return
    cell.value = None


# ---------------------------------------------------------------------------
# Feuille DIAMETRE
# ---------------------------------------------------------------------------

_COLS_DIAMETRE_A_VIDER = {'DMOYMM', 'SMM2', 'DMOYGLOBALMM'}


def _vider_diametre(ws) -> int:
    """Vide D_Moy, S_ et D_Moy Global pour chaque groupe."""
    ligne_entete = None
    for r in (1, 2, 3):
        for c in range(1, ws.max_column + 1):
            if _norm(ws.cell(r, c).value) == 'D1MM':
                ligne_entete = r
                break
        if ligne_entete:
            break

    if ligne_entete is None:
        return 0

    colonnes = set()
    groupes = 0
    for c in range(1, ws.max_column + 1):
        entete = _norm(ws.cell(ligne_entete, c).value)
        if entete in _COLS_DIAMETRE_A_VIDER:
            colonnes.add(c)
            if entete == 'DMOYMM':
                groupes += 1

    for col in colonnes:
        for r in range(ligne_entete + 1, ws.max_row + 1):
            _effacer(ws, r, col)

    return groupes


# ---------------------------------------------------------------------------
# Feuilles COROUI / COSOU
# ---------------------------------------------------------------------------

def _vider_courbe(ws) -> int:
    """Vide Def, Cont, Cont(max), Def rupt, E pour chaque bloc Elong/Force."""
    ligne_entete = 1  # convention observée sur les fichiers CR/CS

    colonnes = set()
    blocs = 0
    for c in range(1, ws.max_column + 1):
        if _norm(ws.cell(ligne_entete, c).value).startswith('ELONG'):
            blocs += 1
            # offsets +2..+6 depuis Elong : Def, Cont, Cont(max), Def rupt, E
            for offset in range(2, 7):
                if c + offset <= ws.max_column:
                    colonnes.add(c + offset)

    for col in colonnes:
        for r in range(ligne_entete + 1, ws.max_row + 1):
            _effacer(ws, r, col)

    return blocs


# ---------------------------------------------------------------------------
# Feuille RECAP
# ---------------------------------------------------------------------------

def _vider_recap(ws) -> int:
    n = 0
    for r in range(1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            if ws.cell(r, c).value is not None:
                _effacer(ws, r, c)
                n += 1
    try:
        ws._charts.clear()
    except Exception:
        pass
    return n


# ---------------------------------------------------------------------------
# Pipeline
# ---------------------------------------------------------------------------

def vider_classeur(source, destination=None, *, supprimer_recap=False, verbeux=True):
    src = Path(source)
    if not src.exists():
        raise FileNotFoundError(f"Fichier introuvable : {src}")

    if destination is None:
        destination = src.with_name(f"{src.stem}_BRUT{src.suffix}")
    dst = Path(destination)
    dst.parent.mkdir(parents=True, exist_ok=True)

    wb = load_workbook(src, data_only=False)

    feuilles_a_supprimer = []
    for nom in list(wb.sheetnames):
        ws = wb[nom]
        genre = _classer(nom)

        if genre == 'diametre':
            g = _vider_diametre(ws)
            if verbeux:
                print(f"  [{nom}] {g} groupe(s) de diamètres vidé(s)")
        elif genre == 'courbe':
            b = _vider_courbe(ws)
            if verbeux:
                print(f"  [{nom}] {b} bloc(s) d'échantillon vidé(s)")
            try:
                ws._charts.clear()
            except Exception:
                pass
        elif genre == 'recap':
            if supprimer_recap:
                feuilles_a_supprimer.append(nom)
                if verbeux:
                    print(f"  [{nom}] marquée pour suppression")
            else:
                n = _vider_recap(ws)
                if verbeux:
                    print(f"  [{nom}] {n} cellule(s) vidée(s)")

    for nom in feuilles_a_supprimer:
        wb.remove(wb[nom])

    wb.save(dst)
    if verbeux:
        print(f"\n→ Fichier brut écrit : {dst.resolve()}")
    return dst


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------
def main(argv=None):
    p = argparse.ArgumentParser()
    p.add_argument("sources", nargs="+")
    p.add_argument("-o", "--sortie")
    p.add_argument("-d", "--dossier-sortie")
    p.add_argument("--supprimer-recap", action="store_true")
    p.add_argument("-q", "--quiet", action="store_true")
    a = p.parse_args(argv)

    # Expansion des jokers côté script (utile sous PowerShell)
    sources_etendues = []
    for s in a.sources:
        s_path = Path(s)
        if any(ch in s for ch in "*?["):
            sources_etendues.extend(sorted(s_path.parent.glob(s_path.name)))
        else:
            sources_etendues.append(s_path)

    if a.sortie and len(sources_etendues) > 1:
        p.error("-o une seule source")
    if a.sortie and a.dossier_sortie:
        p.error("-o ou -d, pas les deux")

    for src in sources_etendues:
        if not src.exists():
            print(f"Introuvable : {src}")
            continue
        if a.sortie:
            dst = a.sortie
        elif a.dossier_sortie:
            dst = Path(a.dossier_sortie) / f"{src.stem}_BRUT{src.suffix}"
        else:
            dst = None
        if not a.quiet:
            print(f"Traitement : {src}")
        vider_classeur(src, dst,
                       supprimer_recap=a.supprimer_recap,
                       verbeux=not a.quiet)
    return 0