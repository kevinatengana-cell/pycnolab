"""
Script de test autonome pour reader.py — version multi-formes.

Usage :
    python test_reader.py
    python test_reader.py --fichier-avec-dim chemin.xlsx
    python test_reader.py --fichier-sans-dim chemin.xlsx

Si aucun chemin n'est fourni, le script cherche dans le dossier courant
des fichiers aux noms probables, et signale ceux qui manquent sans planter.
"""

from __future__ import annotations

import argparse
import sys
import traceback
from pathlib import Path

# --- Import du module à tester ------------------------------------------------
try:
    from reader import (
        read_workbook,
        lire_donnees_brutes,
        _lire_dimensions_pour_forme,
        _lire_diametres_generiques,
        _trouver_feuille_dimensions,
        _FORMES_VALIDES,
    )
except ImportError as e:
    print(f"[FATAL] Impossible d'importer reader.py : {e}")
    print("        Lance ce script depuis le dossier contenant reader.py.")
    sys.exit(1)

from openpyxl import load_workbook


# --- Utilitaires d'affichage --------------------------------------------------

VERT = "\033[92m"
ROUGE = "\033[91m"
JAUNE = "\033[93m"
GRIS = "\033[90m"
RESET = "\033[0m"
GRAS = "\033[1m"

_resultats = []  # (nom_test, statut, message)


def _affiche_titre(titre: str):
    print()
    print(f"{GRAS}{'─' * 70}{RESET}")
    print(f"{GRAS}  {titre}{RESET}")
    print(f"{GRAS}{'─' * 70}{RESET}")


def _ok(nom: str, message: str = ""):
    _resultats.append((nom, "OK", message))
    print(f"  {VERT}[OK]{RESET}    {nom}" + (f"  {GRIS}{message}{RESET}" if message else ""))


def _echec(nom: str, message: str):
    _resultats.append((nom, "ECHEC", message))
    print(f"  {ROUGE}[ECHEC]{RESET} {nom}")
    print(f"          {ROUGE}{message}{RESET}")


def _skip(nom: str, raison: str):
    _resultats.append((nom, "SKIP", raison))
    print(f"  {JAUNE}[SKIP]{RESET}  {nom}  {GRIS}({raison}){RESET}")


def _info(message: str):
    print(f"  {GRIS}{message}{RESET}")


# --- Détection des fichiers de test -------------------------------------------

NOMS_POSSIBLES_AVEC_DIM = [
    "IMPORT_TRACTION.xlsx",
    "EXEMPLE_ENTREE_TRACTION_MULTI_FORMES.xlsx",
    "EXEMPLE_ENTREE_COMPRESSION_MULTI_FORMES.xlsx",
    "EXEMPLE_ENTREE_FLEXION_3_POINTS_MULTI_FORMES.xlsx",
]

NOMS_POSSIBLES_SANS_DIM = [
    "SANS_DIMENSIONS.xlsx",
    "IMPORT_SANS_DIM.xlsx",
]


def _chercher_fichier(noms: list[str]) -> Path | None:
    for nom in noms:
        p = Path(nom)
        if p.exists():
            return p
    return None


# --- Tests --------------------------------------------------------------------

def test_01_detection_feuille(chemin: Path):
    """La feuille dimensions est-elle bien détectée ?"""
    _affiche_titre("Test 1 — Détection de la feuille de dimensions")
    try:
        wb = load_workbook(chemin, data_only=True, read_only=True)
        try:
            nom = _trouver_feuille_dimensions(wb)
            if nom is None:
                _echec("_trouver_feuille_dimensions", "Aucune feuille détectée")
                return
            _ok("_trouver_feuille_dimensions", f"feuille trouvée : {nom!r}")
        finally:
            wb.close()
    except Exception as e:
        _echec("_trouver_feuille_dimensions", f"{type(e).__name__} : {e}")


def test_02_lecture_donnees_brutes(chemin: Path):
    """Le flag a_feuille_dimensions remonte bien ?"""
    _affiche_titre("Test 2 — Lecture brute + flag a_feuille_dimensions")
    try:
        info = lire_donnees_brutes(str(chemin), "traction")
        flag = info.get("a_feuille_dimensions")
        gammes = info.get("gammes", [])
        if not isinstance(flag, bool):
            _echec("a_feuille_dimensions", f"attendu bool, reçu {type(flag).__name__}")
            return
        _ok("a_feuille_dimensions", f"= {flag}")
        if not gammes:
            _echec("gammes", "aucune gamme détectée")
            return
        _ok("gammes", f"{len(gammes)} gamme(s) : {[g['nom_gamme'] for g in gammes]}")
        for g in gammes:
            _info(f"  - {g['nom_gamme']} : {len(g['echantillons'])} échantillon(s)")
    except Exception as e:
        _echec("lire_donnees_brutes", f"{type(e).__name__} : {e}")
        traceback.print_exc()


def test_03_lecture_forme(chemin: Path, forme: str):
    """Lecture d'une forme précise."""
    _affiche_titre(f"Test 3 — Lecture de la forme {forme!r}")
    try:
        wb = load_workbook(chemin, data_only=True, read_only=True)
        try:
            dims = _lire_dimensions_pour_forme(wb, forme)
        finally:
            wb.close()

        if not dims:
            _skip(f"lecture {forme}", "aucun bloc trouvé")
            return

        _ok(f"lecture {forme}", f"{len(dims)} bloc(s)")
        for nom_gamme, echantillons in dims.items():
            _info(f"  - {nom_gamme} : {len(echantillons)} éch.")
            for ech in echantillons[:2]:
                _info(
                    f"      {ech['identifiant']}  "
                    f"forme={ech['forme']}  "
                    f"S={ech['section_mm2']:.4f} mm²"
                )
    except Exception as e:
        _echec(f"lecture {forme}", f"{type(e).__name__} : {e}")
        traceback.print_exc()


def test_04_read_workbook_avec_dim(chemin: Path, forme: str):
    """read_workbook doit remplir les sections depuis le fichier."""
    _affiche_titre(f"Test 4 — read_workbook avec feuille (forme={forme!r})")
    try:
        res = read_workbook(str(chemin), "traction", forme=forme)
        total = 0
        for g in res["gammes"]:
            for e in g["echantillons"]:
                if "section_mm2" not in e or e["section_mm2"] <= 0:
                    _echec(
                        f"section de {e['identifiant']}",
                        f"section manquante ou nulle : {e.get('section_mm2')}",
                    )
                    return
                total += 1
        _ok("read_workbook", f"{total} échantillon(s) avec section valide")
        for g in res["gammes"]:
            e = g["echantillons"][0]
            _info(
                f"  - {g['nom_gamme']} : "
                f"{e['identifiant']}  forme={e.get('forme')}  "
                f"S={e['section_mm2']:.4f}"
            )
    except Exception as e:
        _echec("read_workbook", f"{type(e).__name__} : {e}")
        traceback.print_exc()


def test_05_mauvais_forme(chemin: Path, bonne_forme: str):
    """Mauvais choix de forme -> erreur claire."""
    _affiche_titre("Test 5 — Mauvais choix de forme (doit lever une erreur)")
    formes_possibles = [f for f in _FORMES_VALIDES if f != bonne_forme]
    mauvaise = formes_possibles[0] if formes_possibles else None
    if mauvaise is None:
        _skip("mauvais choix de forme", "aucune autre forme disponible")
        return
    try:
        read_workbook(str(chemin), "traction", forme=mauvaise)
        _echec(
            "erreur attendue",
            f"aucune erreur levée pour forme={mauvaise!r}",
        )
    except ValueError as e:
        msg = str(e)
        if "correspondant" in msg or "introuvable" in msg or "Aucune" in msg:
            _ok("erreur claire levée", f"{msg[:80]}…")
        else:
            _echec("erreur levée mais message inattendu", msg)
    except Exception as e:
        _echec("mauvaise exception", f"{type(e).__name__} : {e}")


def test_06_sans_feuille_avec_geometrie(chemin: Path):
    """Fichier sans feuille + géométrie fournie -> OK."""
    _affiche_titre("Test 6 — Fichier sans feuille + géométrie UI")
    try:
        res = read_workbook(
            str(chemin),
            "traction",
            geometrie={"section_mm2": 40.0, "l0_mm": 20.0},
            forme="rectangle",
        )
        total = 0
        for g in res["gammes"]:
            for e in g["echantillons"]:
                if abs(e["section_mm2"] - 40.0) > 1e-9:
                    _echec(
                        f"section de {e['identifiant']}",
                        f"attendu 40.0, reçu {e['section_mm2']}",
                    )
                    return
                total += 1
        _ok("sections injectées", f"{total} échantillon(s) à 40.0 mm²")
    except Exception as e:
        _echec("read_workbook sans feuille", f"{type(e).__name__} : {e}")
        traceback.print_exc()


def test_07_sans_feuille_sans_geometrie(chemin: Path):
    """Fichier sans feuille + pas de géométrie -> erreur claire."""
    _affiche_titre("Test 7 — Fichier sans feuille + pas de géométrie (doit lever)")
    try:
        read_workbook(str(chemin), "traction", forme="rectangle")
        _echec("erreur attendue", "aucune erreur levée")
    except ValueError as e:
        msg = str(e)
        if "Pas de feuille" in msg or "section saisie" in msg:
            _ok("erreur claire levée", f"{msg[:80]}…")
        else:
            _echec("message inattendu", msg)
    except Exception as e:
        _echec("mauvaise exception", f"{type(e).__name__} : {e}")


def test_08_retrocompatibilite(chemin: Path):
    """L'alias _lire_diametres_generiques renvoie toujours des diamètres."""
    _affiche_titre("Test 8 — Rétrocompatibilité _lire_diametres_generiques")
    try:
        wb = load_workbook(chemin, data_only=True, read_only=True)
        try:
            dims = _lire_diametres_generiques(wb)
        finally:
            wb.close()

        if not dims:
            _skip("_lire_diametres_generiques", "aucun bloc cylindre trouvé")
            return
        _ok("_lire_diametres_generiques", f"{len(dims)} bloc(s)")
        for nom_gamme, echantillons in dims.items():
            for e in echantillons[:1]:
                if "diametres_mm" not in e:
                    _echec(
                        f"diametres_mm de {e['identifiant']}",
                        "clé manquante dans le retour",
                    )
                    return
        _ok("clé diametres_mm présente", "compat export assurée")
    except Exception as e:
        _echec("_lire_diametres_generiques", f"{type(e).__name__} : {e}")
        traceback.print_exc()


# --- Programme principal ------------------------------------------------------

def _resume():
    _affiche_titre("RÉSUMÉ")
    total = len(_resultats)
    ok = sum(1 for _, s, _ in _resultats if s == "OK")
    echecs = sum(1 for _, s, _ in _resultats if s == "ECHEC")
    skips = sum(1 for _, s, _ in _resultats if s == "SKIP")

    print(f"  Total   : {total}")
    print(f"  {VERT}OK{RESET}      : {ok}")
    print(f"  {ROUGE}ECHEC{RESET}   : {echecs}")
    print(f"  {JAUNE}SKIP{RESET}    : {skips}")
    print()
    if echecs:
        print(f"  {ROUGE}Tests en échec :{RESET}")
        for nom, s, msg in _resultats:
            if s == "ECHEC":
                print(f"    - {nom}  {GRIS}({msg}){RESET}")
        print()
        print(f"{ROUGE}{GRAS}ÉCHEC GLOBAL{RESET}")
        return 1
    print(f"{VERT}{GRAS}TOUS LES TESTS PASSENT{RESET}")
    return 0


def main():
    parser = argparse.ArgumentParser(description="Tests de reader.py (multi-formes)")
    parser.add_argument(
        "--fichier-avec-dim",
        type=str,
        default=None,
        help="Classeur contenant une feuille de dimensions",
    )
    parser.add_argument(
        "--fichier-sans-dim",
        type=str,
        default=None,
        help="Classeur SANS feuille de dimensions (une gamme minimale suffit)",
    )
    parser.add_argument(
        "--forme",
        type=str,
        default="cylindre",
        help="Forme à tester pour la lecture (défaut : cylindre)",
    )
    args = parser.parse_args()

    # Résolution du fichier "avec dimensions"
    if args.fichier_avec_dim:
        chemin_avec = Path(args.fichier_avec_dim)
        if not chemin_avec.exists():
            print(f"{ROUGE}Introuvable : {chemin_avec}{RESET}")
            return 2
    else:
        chemin_avec = _chercher_fichier(NOMS_POSSIBLES_AVEC_DIM)
        if chemin_avec is None:
            print(f"{ROUGE}Aucun fichier de test 'avec dimensions' trouvé.{RESET}")
            print(f"Cherchés : {NOMS_POSSIBLES_AVEC_DIM}")
            print("Utilise --fichier-avec-dim pour préciser.")
            return 2

    chemin_sans = None
    if args.fichier_sans_dim:
        chemin_sans = Path(args.fichier_sans_dim)
        if not chemin_sans.exists():
            print(f"{JAUNE}Chemin 'sans dimensions' introuvable, tests 6 et 7 ignorés.{RESET}")
            chemin_sans = None
    else:
        chemin_sans = _chercher_fichier(NOMS_POSSIBLES_SANS_DIM)

    print(f"{GRAS}Fichier avec dimensions :{RESET} {chemin_avec}")
    if chemin_sans:
        print(f"{GRAS}Fichier sans dimensions :{RESET} {chemin_sans}")
    else:
        print(f"{GRAS}Fichier sans dimensions :{RESET} {JAUNE}non fourni{RESET}")
    print(f"{GRAS}Forme testée           :{RESET} {args.forme}")

    # Tests qui nécessitent un fichier avec feuille
    test_01_detection_feuille(chemin_avec)
    test_02_lecture_donnees_brutes(chemin_avec)
    test_03_lecture_forme(chemin_avec, args.forme)
    test_04_read_workbook_avec_dim(chemin_avec, args.forme)
    test_05_mauvais_forme(chemin_avec, args.forme)
    test_08_retrocompatibilite(chemin_avec)

    # Tests qui nécessitent un fichier sans feuille
    if chemin_sans:
        test_06_sans_feuille_avec_geometrie(chemin_sans)
        test_07_sans_feuille_sans_geometrie(chemin_sans)
    else:
        _affiche_titre("Tests 6 et 7 — Fichier sans feuille")
        _skip("tests 6 et 7", "aucun fichier sans dimensions fourni")

    return _resume()


if __name__ == "__main__":
    sys.exit(main())