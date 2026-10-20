from typing import Dict, Any, Optional, List
import openpyxl
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.chart import ScatterChart, Reference, Series
from openpyxl.utils import column_index_from_string

from moteur_python.calculs.traction import TractionResultats

class ExcelDynamicMapper:
    """Injecte dynamiquement les résultats de calcul dans une feuille Excel 
    sans dépendre d'une mise en page figée.
    """

    def __init__(self, mapping_variables: Dict[str, str]):
        """
        :param mapping_variables: Dictionnaire associant une variable à sa cellule.
                                  Exemple transmis par Flutter :
                                  {"E": "C10", "Rm": "C11", "Rp02": "C12", "A%": "C13"}
        """
        self.mapping = mapping_variables

    def injecter_resultats(self, ws: Worksheet, resultats: TractionResultats) -> None:
        """Parcourt le dictionnaire de correspondance et écrit les valeurs dans la feuille."""
        valeurs_disponibles = {
            "E": resultats.module_young_E_MPa,
            "Rm": resultats.resistance_mecanique_Rm_MPa,
            "Rp02": resultats.limite_elasticite_Rp02_MPa,
            "A%": resultats.allongement_rupture_A_pct,
        }

        for cle_variable, adresse_cellule in self.mapping.items():
            if cle_variable in valeurs_disponibles and adresse_cellule:
                valeur = valeurs_disponibles[cle_variable]
                if valeur is not None:
                    ws[adresse_cellule] = valeur

    def injecter_tableau_points(
        self, 
        ws: Worksheet, 
        colonne_dep: str, 
        colonne_force: str, 
        ligne_depart: int, 
        deplacements: List[float], 
        forces: List[float]
    ) -> None:
        """Injecte les séries temporelles (effort/déplacement) sur la hauteur nécessaire."""
        for i, (d, f) in enumerate(zip(deplacements, forces)):
            num_ligne = ligne_depart + i
            ws[f"{colonne_dep}{num_ligne}"] = d
            ws[f"{colonne_force}{num_ligne}"] = f

    def ajouter_graphique_traction(
        self, 
        ws: Worksheet, 
        colonne_dep: str, 
        colonne_force: str, 
        ligne_debut: int, 
        nb_points: int, 
        cellule_ancrage: str = "E2",
        titre: str = "Courbe Effort / Déplacement"
    ) -> None:
        """Génère un graphique ScatterChart dynamique ancré à la cellule spécifiée."""
        chart = ScatterChart()
        chart.title = titre
        chart.style = 13
        chart.x_axis.title = 'Déplacement (mm)'
        chart.y_axis.title = 'Force (N)'

        col_dep_idx = column_index_from_string(colonne_dep)
        col_force_idx = column_index_from_string(colonne_force)

        xvalues = Reference(ws, min_col=col_dep_idx, min_row=ligne_debut, max_row=ligne_debut + nb_points - 1)
        yvalues = Reference(ws, min_col=col_force_idx, min_row=ligne_debut, max_row=ligne_debut + nb_points - 1)
        
        series = Series(yvalues, xvalues, title_from_data=False)
        series.graphicalProperties.line.solidFill = "0000FF"
        series.graphicalProperties.line.width = 25000  # Épaisseur de ligne
        
        chart.series.append(series)
        chart.width = 15
        chart.height = 10
        
        ws.add_chart(chart, cellule_ancrage)