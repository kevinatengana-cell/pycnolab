import asyncio
import tempfile
import unittest
import zipfile
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.utils.units import cm_to_EMU

from moteur_python.api.main import (
    DonneesEssaiRequest,
    EssaiExcelRequest,
    app,
    health,
    route_calculer_depuis_excel,
    route_calculer_donnees_capture,
    route_exporter_essai_excel,
    route_lister_echantillons,
    route_lire_donnees_brutes,
)
from moteur_python.export_excel.production_workbook_exporter import (
    _blocks,
    _ensure_block,
    _largeur_col_cm,
    _young,
)
from moteur_python.import_excel.reader import (
    lire_donnees_brutes,
    lire_echantillons_depuis_excel,
)


class EssaiExcelApiTests(unittest.TestCase):
    def setUp(self):
        self.dossier = tempfile.TemporaryDirectory()
        self.fichier_source = Path(self.dossier.name) / 'compression.xlsx'

        workbook = Workbook()
        feuille = workbook.active
        feuille.title = 'F0'
        feuille.cell(1, 1, 'Elong.(mm) - ECH 1')
        feuille.cell(1, 2, 'Force(N) - ECH 1')
        for index in range(1, 26):
            feuille.cell(index + 1, 1, index * 0.1)
            feuille.cell(index + 1, 2, index * 10)
        workbook.save(self.fichier_source)

    def tearDown(self):
        self.dossier.cleanup()

    def test_routes_sont_publiees_et_protocole_incrementee(self):
        paths = app.openapi()['paths']
        self.assertIn('/essais/{type_essai}/lister-echantillons', paths)
        self.assertIn('/essais/{type_essai}/lire-brut', paths)
        self.assertIn('/essais/{type_essai}/calculer', paths)
        self.assertIn('/essais/{type_essai}/calculer-depuis-excel', paths)
        self.assertIn('/essais/{type_essai}/exporter-excel', paths)
        self.assertEqual(asyncio.run(health()), {'status': 'ok'})

    def test_lister_echantillons_ne_requiert_pas_de_geometrie(self):
        resultat = route_lister_echantillons(
            'compression',
            EssaiExcelRequest(chemin_fichier=str(self.fichier_source)),
        )
        self.assertEqual(resultat['gammes'], [{'nom_gamme': 'F0', 'echantillons': ['ECH1']}])

    def test_lire_brut_retourne_points_sans_geometrie(self):
        resultat = route_lire_donnees_brutes(
            'compression',
            EssaiExcelRequest(chemin_fichier=str(self.fichier_source)),
        )
        self.assertEqual(resultat['type_essai'], 'compression')
        self.assertEqual(
            resultat['gammes'][0]['echantillons'][0]['points_courbe'][0],
            {'deplacement_mm': 0.1, 'force_newton': 10.0},
        )

    def test_identifiants_diametre_plot_casse_et_espaces_matchent_ech_positionnel(self):
        fichier = Path(self.dossier.name) / 'traction_diametres.xlsx'
        workbook = Workbook()
        diametres = workbook.active
        diametres.title = 'DIAMETRE'
        diametres.cell(1, 2, 'F0')
        for column, label in enumerate(('D1', 'D2', 'D3'), start=2):
            diametres.cell(2, column, label)
        diametres.cell(3, 1, 'plot 1')
        diametres.cell(4, 1, 'eCh 2')
        for column, value in enumerate((2.0, 2.1, 1.9), start=2):
            diametres.cell(3, column, value)
        for column, value in enumerate((3.0, 3.1, 2.9), start=2):
            diametres.cell(4, column, value)

        courbe = workbook.create_sheet('F0')
        courbe.cell(1, 1, 'Elong.(mm) - PLOT 1')
        courbe.cell(1, 2, 'Force(N) - Ech 1')
        courbe.cell(1, 3, 'ELONG.(mm) - eCH 2')
        courbe.cell(1, 4, 'force(N) - PLOT 2')
        courbe.cell(2, 1, 0.1)
        courbe.cell(2, 2, 10)
        courbe.cell(2, 3, 0.2)
        courbe.cell(2, 4, 20)
        workbook.save(fichier)
        workbook.close()

        resultat = lire_donnees_brutes(fichier, 'traction')
        echantillons = resultat['gammes'][0]['echantillons']
        self.assertEqual([ech['identifiant'] for ech in echantillons], ['ECH1', 'ECH2'])
        self.assertAlmostEqual(echantillons[0]['diametre_moyen_mm'], 2.0)
        self.assertAlmostEqual(echantillons[1]['diametre_moyen_mm'], 3.0)
        self.assertEqual(len(echantillons[0]['points_courbe']), 1)

        legacy_samples = lire_echantillons_depuis_excel(str(fichier))
        self.assertEqual(
            [sample['identifiant'] for sample in legacy_samples],
            ['ECH1', 'ECH2'],
        )
        self.assertNotIn('diametre_cr', legacy_samples[0])
        self.assertNotIn('diametre_cs', legacy_samples[0])

    def test_entetes_courbe_sont_insensibles_a_la_casse_et_au_prefixe_plot(self):
        fichier = Path(self.dossier.name) / 'courbes_variantes.xlsx'
        workbook = Workbook()
        feuille = workbook.active
        feuille.title = 'F0'
        headers = (
            'eLoNg.(mm) - pLoT 1', 'fOrCe(N) - PLOT 1',
            'ELONG.(mm) - eCh 2', 'FORCE(N) - eCH 2',
            'Elong.(mm) - ECH 3', 'Force(N) - ech 3',
        )
        for column, header in enumerate(headers, start=1):
            feuille.cell(1, column, header)
        for column in (1, 3, 5):
            feuille.cell(2, column, 0.1)
            feuille.cell(2, column + 1, 10)
        workbook.save(fichier)
        workbook.close()

        resultat = lire_donnees_brutes(fichier, 'compression')
        self.assertEqual(
            [sample['identifiant'] for sample in resultat['gammes'][0]['echantillons']],
            ['ECH1', 'ECH2', 'ECH3'],
        )

    def test_exporteur_reutilise_un_bloc_plot_sans_en_creer_un_second(self):
        workbook = Workbook()
        feuille = workbook.active
        feuille.cell(1, 1, 'Elong.(mm) - PLOT 1')
        feuille.cell(1, 2, 'Force(N) - Plot 1')

        self.assertEqual(_blocks(feuille), {1: (1, 2)})
        self.assertEqual(_ensure_block(feuille, 1), (1, 2))
        self.assertEqual(feuille.max_column, 2)
        self.assertEqual(feuille.cell(1, 1).value, 'Elong.(mm) - PLOT 1')
        workbook.close()

    def test_calculer_depuis_excel_retourne_les_resultats(self):
        resultat = route_calculer_depuis_excel(
            'compression',
            EssaiExcelRequest(
                chemin_fichier=str(self.fichier_source),
                geometrie_compression={'l0_mm': 20, 'section_mm2': 10},
            ),
        )
        self.assertEqual(resultat['type_essai'], 'compression')
        self.assertEqual(resultat['gammes'][0]['nom_gamme'], 'F0')
        self.assertEqual(resultat['gammes'][0]['echantillons'][0]['identifiant'], 'ECH1')

    def test_recalculer_depuis_les_donnees_stockees_sans_fichier(self):
        resultat = route_calculer_donnees_capture(
            'compression',
            DonneesEssaiRequest(
                donnees={
                    'gammes': [{
                        'nom_gamme': 'F0',
                        'echantillons': [{
                            'identifiant': 'ECH1',
                            'l0_mm': 20.0,
                            'section_mm2': 10.0,
                            'points_courbe': [
                                {'deplacement_mm': index * 0.1, 'force_newton': index * 10.0}
                                for index in range(1, 26)
                            ],
                        }],
                    }],
                },
            ),
        )
        self.assertEqual(resultat['type_essai'], 'compression')
        self.assertEqual(resultat['gammes'][0]['echantillons'][0]['identifiant'], 'ECH1')

    def test_exporter_excel_retourne_un_classeur_xlsx(self):
        reponse = route_exporter_essai_excel(
            'compression',
            EssaiExcelRequest(
                chemin_fichier=str(self.fichier_source),
                geometrie_compression={'l0_mm': 20, 'section_mm2': 10},
            ),
        )
        self.assertEqual(reponse.status_code, 200)
        with zipfile.ZipFile(BytesIO(reponse.body)) as classeur:
            self.assertIsNone(classeur.testzip())
        workbook = load_workbook(BytesIO(reponse.body))
        feuille = workbook['F0']
        (elong_colonne, force_colonne), = _blocks(feuille).values()
        largeur_attendue = sum(
            _largeur_col_cm(feuille, colonne)
            for colonne in range(elong_colonne, force_colonne + 6)
        )
        self.assertEqual(
            feuille._charts[0].anchor.ext.cx,
            cm_to_EMU(largeur_attendue),
        )
        workbook.close()

    def test_exporteur_cree_le_graphique_si_jeune_indetermine(self):
        fichier = Path(self.dossier.name) / 'compression_plateau.xlsx'
        workbook = Workbook()
        feuille = workbook.active
        feuille.title = 'F0'
        feuille.cell(1, 1, 'Elong.(mm) - ECH 1')
        feuille.cell(1, 2, 'Force(N) - ECH 1')
        for row, (elong, force) in enumerate(
            ((0.154, 0.067), (0.221, 0.067), (0.221, 0.211)),
            start=2,
        ):
            feuille.cell(row, 1, elong)
            feuille.cell(row, 2, force)
        workbook.save(fichier)
        workbook.close()

        self.assertIsNone(
            _young(
                [0.154 / 20, 0.221 / 20, 0.221 / 20],
                [0.067 / 10, 0.067 / 10, 0.211 / 10],
            )[0]
        )
        reponse = route_exporter_essai_excel(
            'compression',
            EssaiExcelRequest(
                chemin_fichier=str(fichier),
                geometrie_compression={'l0_mm': 20, 'section_mm2': 10},
            ),
        )
        self.assertEqual(reponse.status_code, 200)

        workbook = load_workbook(BytesIO(reponse.body))
        feuille = workbook['F0']
        self.assertEqual(len(feuille._charts), 1)
        serie = feuille._charts[0].series[0]
        self.assertIn('$C$2:$C$4', serie.xVal.numRef.f)
        self.assertIn('$D$2:$D$4', serie.yVal.numRef.f)
        workbook.close()


if __name__ == '__main__':
    unittest.main()