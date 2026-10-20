"""
Génère un template Excel propre pour l'exporteur PycnoLab.

Structure produite :
    DIMENSIONS      → feuille vide, le code y écrit le bloc adapté à la forme
    MODELE_GAMME    → ligne 1 avec les 7 en-têtes de colonnes, styles associés
    RECCAPITULATIF  → feuille vide, reconstruite par le code

Usage :
    python creer_template.py
    python creer_template.py --sortie "chemin/vers/MODELE_TRACTION_VIDE.xlsx"
    python creer_template.py --sortie "..." --titre "DIMENSIONS DES ÉPROUVETTES"
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# --- Noms des feuilles --------------------------------------------------------

FEUILLE_DIMENSIONS = 'DIMENSIONS'
FEUILLE_MODELE_GAMME = 'MODELE_GAMME'
FEUILLE_RECAP = 'RECCAPITULATIF'   # deux C, comme la constante du code


# --- Styles -------------------------------------------------------------------

POLICE = 'Times New Roman'
TAILLE_EN_TETE = 11
TAILLE_TITRE = 12

FONT_TITRE = Font(name=POLICE, size=TAILLE_TITRE, bold=True, color='1F3864')
FONT_EN_TETE = Font(name=POLICE, size=TAILLE_EN_TETE, bold=True, color='1F3864')
FONT_NORMAL = Font(name=POLICE, size=TAILLE_EN_TETE)

FILL_EN_TETE = PatternFill('solid', fgColor='BDD7EE')

BORDURE_FINE = Border(
    left=Side(style='thin', color='B7B7B7'),
    right=Side(style='thin', color='B7B7B7'),
    top=Side(style='thin', color='B7B7B7'),
    bottom=Side(style='thin', color='B7B7B7'),
)

ALIGN_CENTRE = Alignment(horizontal='center', vertical='center', wrap_text=True)
ALIGN_GAUCHE = Alignment(horizontal='left', vertical='center')


# --- En-têtes de la feuille modèle de gamme -----------------------------------

EN_TETES_GAMME = [
    'Elong.(mm) - Ech 1',
    'Force(N) - Ech 1',
    'Def',
    'Cont',
    'Cont(max)',
    'Def rupt',
    'E(mpa)',
]

LARGEURS_GAMME = [18, 18, 12, 12, 12, 12, 12]


# --- Construction -------------------------------------------------------------

def _construire_workbook(titre_dimensions: str | None) -> Workbook:
    wb = Workbook()

    # 1. Feuille DIMENSIONS (vide, avec éventuellement un titre)
    ws_dim = wb.active
    ws_dim.title = FEUILLE_DIMENSIONS
    if titre_dimensions:
        cell = ws_dim.cell(1, 1, titre_dimensions)
        cell.font = FONT_TITRE
        cell.alignment = ALIGN_GAUCHE
    ws_dim.column_dimensions['A'].width = 14
    for col in range(2, 12):
        ws_dim.column_dimensions[get_column_letter(col)].width = 12

    # 2. Feuille MODELE_GAMME
    ws_gamme = wb.create_sheet(FEUILLE_MODELE_GAMME)
    for j, entete in enumerate(EN_TETES_GAMME, start=1):
        cell = ws_gamme.cell(1, j, entete)
        cell.font = FONT_EN_TETE
        cell.fill = FILL_EN_TETE
        cell.alignment = ALIGN_CENTRE
        cell.border = BORDURE_FINE
    for j, largeur in enumerate(LARGEURS_GAMME, start=1):
        ws_gamme.column_dimensions[get_column_letter(j)].width = largeur
    # Hauteur de la ligne d'en-tête un peu plus grande
    ws_gamme.row_dimensions[1].height = 28
    # Figer la ligne d'en-tête (confort visuel dans le rapport final)
    ws_gamme.freeze_panes = 'A2'

    # 3. Feuille RECCAPITULATIF (vide)
    ws_recap = wb.create_sheet(FEUILLE_RECAP)
    ws_recap.column_dimensions['A'].width = 22
    for col in range(2, 10):
        ws_recap.column_dimensions[get_column_letter(col)].width = 16

    return wb


def _verifier_structure(wb: Workbook) -> list[str]:
    """Renvoie la liste des avertissements éventuels."""
    avertissements = []

    if FEUILLE_DIMENSIONS not in wb.sheetnames:
        avertissements.append(f"Feuille {FEUILLE_DIMENSIONS!r} absente.")
    if FEUILLE_MODELE_GAMME not in wb.sheetnames:
        avertissements.append(f"Feuille {FEUILLE_MODELE_GAMME!r} absente.")
    if FEUILLE_RECAP not in wb.sheetnames:
        avertissements.append(f"Feuille {FEUILLE_RECAP!r} absente.")

    # Vérifier que la feuille DIMENSIONS est bien vide
    ws = wb[FEUILLE_DIMENSIONS]
    for row in ws.iter_rows(min_row=2, max_row=5, max_col=10):
        for cell in row:
            if cell.value is not None:
                avertissements.append(
                    f"Feuille {FEUILLE_DIMENSIONS!r} : cellule "
                    f"{cell.coordinate} non vide ({cell.value!r})"
                )
                break

    # Vérifier que la feuille RECAP est vide
    ws = wb[FEUILLE_RECAP]
    for row in ws.iter_rows(min_row=1, max_row=5, max_col=10):
        for cell in row:
            if cell.value is not None:
                avertissements.append(
                    f"Feuille {FEUILLE_RECAP!r} : cellule "
                    f"{cell.coordinate} non vide ({cell.value!r})"
                )
                break

    return avertissements


# --- Point d'entrée -----------------------------------------------------------

def main() -> int:
    parser = argparse.ArgumentParser(
        description="Génère un template Excel propre pour PycnoLab."
    )
    parser.add_argument(
        '--sortie', type=str, default=None,
        help='Chemin du fichier à écrire (défaut : MODELE_TRACTION_VIDE.xlsx '
             'à côté de ce script)',
    )
    parser.add_argument(
        '--titre', type=str, default=None,
        help='Titre à mettre en A1 de la feuille DIMENSIONS '
             '(défaut : aucun titre)',
    )
    parser.add_argument(
        '--force', action='store_true',
        help='Écrase le fichier de sortie s\'il existe déjà.',
    )
    args = parser.parse_args()

    if args.sortie:
        sortie = Path(args.sortie)
    else:
        sortie = Path(__file__).parent / 'MODELE_TRACTION_VIDE.xlsx'

    if sortie.exists() and not args.force:
        print(f"Le fichier existe déjà : {sortie}")
        print("Utilise --force pour l'écraser.")
        return 1

    sortie.parent.mkdir(parents=True, exist_ok=True)

    wb = _construire_workbook(args.titre)
    avertissements = _verifier_structure(wb)

    if avertissements:
        print("Avertissements :")
        for a in avertissements:
            print(f"  - {a}")
        print()

    wb.save(sortie)

    print(f"Template généré : {sortie.resolve()}")
    print(f"  Feuilles : {wb.sheetnames}")
    print(f"  Feuille de dimensions : {FEUILLE_DIMENSIONS!r}")
    print(f"  Feuille modèle de gamme : {FEUILLE_MODELE_GAMME!r}")
    print(f"  Feuille récap : {FEUILLE_RECAP!r}")
    return 0


if __name__ == '__main__':
    sys.exit(main())