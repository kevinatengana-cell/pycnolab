from pathlib import Path
import re
import math
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


def norm(v):
    return re.sub(r'[^A-Z0-9]', '', str(v or '').upper())


COLONNES_CALCULEES = {
    'DEFMM', 'DEF', 'CONTMPA', 'CONT', 'CONTMAXMPA', 'CONTMAX',
    'DEFRUPT', 'EMPA', 'DMOYMM', 'SMM2', 'DMOYGLOBALMM',
}
PREFIXES_BRUTS = ('ELONG', 'FORCE', 'D1', 'D2', 'D3')


def effacer(ws, r, c):
    try:
        ws.cell(r, c).value = None
    except AttributeError:
        pass


def trouver_ligne_entete(ws):
    for r in (1, 2):
        for c in range(1, min(ws.max_column, 30) + 1):
            if norm(ws.cell(r, c).value).startswith(PREFIXES_BRUTS):
                return r
    return 1


def extraire_sections(ws):
    """Lit la section dans les formules CONT (=Force/section)."""
    ligne = trouver_ligne_entete(ws)
    sections = []
    for c in range(1, ws.max_column + 1):
        h = norm(ws.cell(ligne, c).value)
        if h in ('CONTMPA', 'CONT'):
            v = ws.cell(ligne + 1, c).value
            if isinstance(v, str) and v.startswith('='):
                m = re.search(r'/\s*\(?\s*([0-9]*\.?[0-9]+)', v)
                if m:
                    try:
                        sections.append(float(m.group(1)))
                        continue
                    except ValueError:
                        pass
            sections.append(None)
    return sections


def vider_colonnes(ws):
    ligne = trouver_ligne_entete(ws)
    cols = []
    for c in range(1, ws.max_column + 1):
        if norm(ws.cell(ligne, c).value) in COLONNES_CALCULEES:
            for r in range(ligne + 1, ws.max_row + 1):
                effacer(ws, r, c)
            cols.append(c)
    return cols


def vider_tout(ws):
    for r in range(1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            effacer(ws, r, c)


def creer_feuille_diametre(wb, sections_par_exp):
    """Recrée une feuille DIAMETRE à partir des sections extraites."""
    for nom in list(wb.sheetnames):
        if norm(nom).startswith('DIAMETRE'):
            del wb[nom]
    ws = wb.create_sheet('DIAMETRE')

    col = 2
    for nom_exp, sections in sections_par_exp.items():
        ws.cell(1, col, nom_exp)
        ws.cell(2, col, 'D_1(mm)')
        ws.cell(2, col + 1, 'D_2(mm)')
        ws.cell(2, col + 2, 'D_3(mm)')

        for i, section in enumerate(sections):
            if section is None or section <= 0:
                continue
            d = 2.0 * math.sqrt(section / math.pi)
            ws.cell(3 + i, col - 1, f'ECH{i + 1}')
            ws.cell(3 + i, col, d)
            ws.cell(3 + i, col + 1, d)
            ws.cell(3 + i, col + 2, d)
        col += 5
    return ws


def traiter(src, dst):
    wb = load_workbook(src)
    sections_par_exp = {}
    rapport = []

    for nom in wb.sheetnames:
        ws = wb[nom]
        n = norm(nom)

        if 'RECAP' in n:
            vider_tout(ws)
            rapport.append((nom, 'VIDEE ENTIEREMENT (recap)', 0))
        else:
            sections = extraire_sections(ws)
            if sections:
                sections_par_exp[nom] = sections
            cols = vider_colonnes(ws)
            rapport.append((nom, f'{len(cols)} col. videe(s)', len(sections)))

    if sections_par_exp:
        creer_feuille_diametre(wb, sections_par_exp)

    wb.save(dst)
    return rapport


def main():
    out_dir = Path('essais/bruts')
    out_dir.mkdir(exist_ok=True)
    fichiers = sorted(Path('essais').glob('*.xlsx'))
    print(f'{len(fichiers)} fichier(s) a traiter\n')

    for f in fichiers:
        dst = out_dir / (f.stem + '_BRUT.xlsx')
        print(f'=== {f.name} ===')
        try:
            rapport = traiter(f, dst)
            for nom, info, n_sec in rapport:
                extra = f' [{n_sec} section(s)]' if n_sec else ''
                print(f'    [{nom}] {info}{extra}')
            print(f'    -> ecrit : {dst.name}\n')
        except Exception as e:
            print(f'    -> ERREUR : {e}\n')


if __name__ == '__main__':
    main()