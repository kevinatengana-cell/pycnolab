"""
Tests de production_workbook_exporter.py — version multi-formes.

Structure attendue :
    pycnolab/
        pycno_core/moteur_python/import_excel/reader.py
        pycno_core/moteur_python/export_excel/production_workbook_exporter.py
        essais/bruts/*.xlsx
        templates/*.xlsx             ← template utilisé pour l'export
        test_exporter.py

Usage :
    python test_exporter.py
    python test_exporter.py --template "chemin/vers/template.xlsx"
    python test_exporter.py --avec-dim "essais/bruts/xxx.xlsx" --forme cylindre
    python test_exporter.py --sans-dim "essais/bruts/yyy.xlsx"

Les rapports générés sont écrits dans essais/exports_test/.
"""

from __future__ import annotations

import argparse
import sys
import traceback
from pathlib import Path

# --- Chemins du projet --------------------------------------------------------

RACINE = Path(__file__).parent
sys.path.insert(0, str(RACINE / 'pycno_core'))

DOSSIER_BRUTS = RACINE / 'essais' / 'bruts'
DOSSIER_SORTIE = RACINE / 'essais' / 'exports_test'
DOSSIER_TEMPLATES = [
    RACINE / 'templates',
    RACINE / 'pycno_core' / 'moteur_python' / 'export_excel' / 'templates',
    RACINE / 'pycno_core' / 'moteur_python' / 'templates',
    RACINE / 'assets' / 'templates',
]


# --- Imports du code à tester -------------------------------------------------

try:
    from moteur_python.import_excel.reader import (
        _trouver_feuille_dimensions,
    )
except ImportError as e:
    print(f"[FATAL] Impossible d'importer reader : {e}")
    sys.exit(1)

try:
    from moteur_python.export_excel.production_workbook_exporter import (
        process_workbook,
    )
except ImportError:
    try:
        from moteur_python.export_excel.production_workbook_exporter import (
            process_workbook,
        )
    except ImportError as e:
        print(f"[FATAL] Impossible d'importer process_workbook : {e}")
        print("Cherche production_workbook_exporter.py :")
        for p in RACINE.rglob('production_workbook_exporter.py'):
            print(f"  {p}")
        sys.exit(1)

from openpyxl import load_workbook


# --- Affichage ----------------------------------------------------------------

VERT = "\033[92m"
ROUGE = "\033[91m"
JAUNE = "\033[93m"
GRIS = "\033[90m"
GRAS = "\033[1m"
RESET = "\033[0m"

_resultats: list[tuple[str, str, str]] = []


def _titre(t: str):
    print()
    print(f"{GRAS}{'─' * 70}{RESET}")
    print(f"{GRAS}  {t}{RESET}")
    print(f"{GRAS}{'─' * 70}{RESET}")


def _ok(nom: str, msg: str = ""):
    _resultats.append((nom, "OK", msg))
    print(f"  {VERT}[OK]{RESET}    {nom}" + (f"  {GRIS}{msg}{RESET}" if msg else ""))


def _echec(nom: str, msg: str):
    _resultats.append((nom, "ECHEC", msg))
    print(f"  {ROUGE}[ECHEC]{RESET} {nom}")
    print(f"          {ROUGE}{msg}{RESET}")


def _skip(nom: str, raison: str):
    _resultats.append((nom, "SKIP", raison))
    print(f"  {JAUNE}[SKIP]{RESET}  {nom}  {GRIS}({raison}){RESET}")


def _info(msg: str):
    print(f"  {GRIS}{msg}{RESET}")


# --- Découverte des ressources ------------------------------------------------

def _trouver_template() -> Path | None:
    """Cherche un template dans les dossiers connus."""
    for dossier in DOSSIER_TEMPLATES:
        if not dossier.exists():
            continue
        for f in sorted(dossier.glob('*.xlsx')):
            nom = f.name.lower()
            if 'template' in nom or 'modele' in nom or 'modele' in nom:
                return f
        # sinon, prendre le premier .xlsx du dossier
        candidats = sorted(dossier.glob('*.xlsx'))
        if candidats:
            return candidats[0]
    return None


def _a_feuille_dimensions(chemin: Path) -> bool:
    try:
        wb = load_workbook(chemin, data_only=True, read_only=True)
        try:
            return _trouver_feuille_dimensions(wb) is not None
        finally:
            wb.close()
    except Exception:
        return False


def _choisir_fichiers(args) -> tuple[Path | None, Path | None]:
    avec = Path(args.avec_dim) if args.avec_dim else None
    sans = Path(args.sans_dim) if args.sans_dim else None

    if avec and not avec.exists():
        print(f"{ROUGE}Introuvable : {avec}{RESET}")
        sys.exit(2)
    if sans and not sans.exists():
        print(f"{ROUGE}Introuvable : {sans}{RESET}")
        sys.exit(2)

    if avec is None or sans is None:
        if not DOSSIER_BRUTS.exists():
            print(f"{ROUGE}Aucun dossier {DOSSIER_BRUTS}{RESET}")
            sys.exit(2)
        candidats = sorted(DOSSIER_BRUTS.glob('*.xlsx'))
        if not candidats:
            print(f"{ROUGE}Aucun .xlsx dans {DOSSIER_BRUTS}{RESET}")
            sys.exit(2)
        _info(f"Analyse de {len(candidats)} fichier(s) dans {DOSSIER_BRUTS}…")
        for f in candidats:
            a_dim = _a_feuille_dimensions(f)
            if avec is None and a_dim:
                avec = f
            if sans is None and not a_dim:
                sans = f

    return avec, sans


# --- Vérifications d'un rapport -----------------------------------------------

def _verifier_fichier(chemin: Path) -> dict:
    """Charge le rapport et renvoie un dict d'informations."""
    wb = load_workbook(chemin, data_only=False)
    info = {
        'feuilles': list(wb.sheetnames),
        'a_feuille_dimensions': False,
        'nom_feuille_dimensions': None,
        'gammes': [],
    }

    for nom in wb.sheetnames:
        norm = nom.upper().replace(' ', '').replace('_', '')
        if norm.startswith(('DIAMETRE', 'DIAMETER', 'DIMENSION')):
            info['a_feuille_dimensions'] = True
            info['nom_feuille_dimensions'] = nom
            break

    # Une feuille de gamme = toute feuille qui n'est ni récap ni dimensions
    for nom in wb.sheetnames:
        if nom == info['nom_feuille_dimensions']:
            continue
        if 'RECAP' in nom.upper().replace(' ', ''):
            continue
        info['gammes'].append(nom)

    wb.close()
    return info


# --- Tests --------------------------------------------------------------------

def test_01_avec_dimensions(template: Path, fichier: Path, forme: str):
    _titre(f"Test 1 — Export avec feuille de dimensions (forme={forme!r})")
    sortie = DOSSIER_SORTIE / f"export_avec_dim_{forme}.xlsx"
    try:
        process_workbook(
            input_path=str(fichier),
            template_path=str(template),
            output_path=str(sortie),
            type_essai='traction',
            l0_mm=20.0,
            forme=forme,
        )
        _ok("process_workbook", f"fichier écrit : {sortie.name}")
    except Exception as e:
        _echec("process_workbook", f"{type(e).__name__} : {e}")
        traceback.print_exc()
        return

    try:
        info = _verifier_fichier(sortie)
        _info(f"  Feuilles : {info['feuilles']}")

        if not info['a_feuille_dimensions']:
            _echec("feuille de dimensions", "absente dans le rapport")
            return
        _ok("feuille de dimensions", f"{info['nom_feuille_dimensions']!r}")

        if not info['gammes']:
            _echec("feuilles de gammes", "aucune détectée")
            return
        _ok("feuilles de gammes", f"{info['gammes']}")

        # Vérifier qu'au moins une feuille de gamme a des formules
        wb = load_workbook(sortie, data_only=False)
        try:
            trouve_formule = False
            for nom in info['gammes']:
                ws = wb[nom]
                for row in ws.iter_rows(
                    min_row=2, max_row=min(5, ws.max_row),
                    min_col=1, max_col=min(20, ws.max_column),
                ):
                    for cell in row:
                        if isinstance(cell.value, str) and cell.value.startswith('='):
                            trouve_formule = True
                            break
                    if trouve_formule:
                        break
                if trouve_formule:
                    break
            if trouve_formule:
                _ok("formules dans gammes", "au moins une formule détectée")
            else:
                _echec("formules dans gammes", "aucune formule détectée")
        finally:
            wb.close()

    except Exception as e:
        _echec("vérification du rapport", f"{type(e).__name__} : {e}")
        traceback.print_exc()


def test_02_sans_dimensions(template: Path, fichier: Path):
    _titre("Test 2 — Export sans feuille de dimensions (sections en littéral)")
    sortie = DOSSIER_SORTIE / "export_sans_dim.xlsx"
    try:
        process_workbook(
            input_path=str(fichier),
            template_path=str(template),
            output_path=str(sortie),
            type_essai='traction',
            l0_mm=20.0,
            geometrie={'section_mm2': 40.0, 'l0_mm': 20.0},
            forme='rectangle',
        )
        _ok("process_workbook", f"fichier écrit : {sortie.name}")
    except Exception as e:
        _echec("process_workbook", f"{type(e).__name__} : {e}")
        traceback.print_exc()
        return

    try:
        info = _verifier_fichier(sortie)
        _info(f"  Feuilles : {info['feuilles']}")

        if info['a_feuille_dimensions']:
            _echec(
                "feuille de dimensions",
                "présente alors que le fichier d'entrée n'en avait pas",
            )
            return
        _ok("pas de feuille de dimensions", "correct")

        if not info['gammes']:
            _echec("feuilles de gammes", "aucune détectée")
            return
        _ok("feuilles de gammes", f"{info['gammes']}")

    except Exception as e:
        _echec("vérification du rapport", f"{type(e).__name__} : {e}")
        traceback.print_exc()


def test_03_mauvais_forme(template: Path, fichier: Path, bonne_forme: str):
    _titre("Test 3 — Mauvais choix de forme (doit lever)")
    sortie = DOSSIER_SORTIE / "export_erreur.xlsx"
    formes_autres = [
        f for f in ('cylindre', 'rectangle', 'carre',
                    'tube_circulaire', 'tube_rectangulaire', 'surface')
        if f != bonne_forme
    ]
    if not formes_autres:
        _skip("mauvais choix", "aucune autre forme")
        return
    mauvaise = formes_autres[0]
    try:
        process_workbook(
            input_path=str(fichier),
            template_path=str(template),
            output_path=str(sortie),
            type_essai='traction',
            l0_mm=20.0,
            forme=mauvaise,
        )
        _echec(
            "erreur attendue",
            f"aucune erreur levée pour forme={mauvaise!r}",
        )
    except ValueError as e:
        msg = str(e)
        if any(k in msg for k in ('correspondant', 'introuvable', 'Aucune', 'Forme')):
            _ok("erreur claire", f"{msg[:70]}…")
        else:
            _echec("message inattendu", msg)
    except Exception as e:
        _echec("mauvaise exception", f"{type(e).__name__} : {e}")


# --- Résumé -------------------------------------------------------------------

def _resume() -> int:
    _titre("RÉSUMÉ")
    total = len(_resultats)
    ok = sum(1 for _, s, _ in _resultats if s == "OK")
    ko = sum(1 for _, s, _ in _resultats if s == "ECHEC")
    sk = sum(1 for _, s, _ in _resultats if s == "SKIP")
    print(f"  Total : {total}")
    print(f"  {VERT}OK{RESET}    : {ok}")
    print(f"  {ROUGE}ECHEC{RESET} : {ko}")
    print(f"  {JAUNE}SKIP{RESET}  : {sk}")
    if ko:
        print()
        print(f"  {ROUGE}En échec :{RESET}")
        for nom, s, msg in _resultats:
            if s == "ECHEC":
                print(f"    - {nom}  {GRIS}({msg}){RESET}")
        print(f"\n{ROUGE}{GRAS}ÉCHEC GLOBAL{RESET}")
        return 1
    print(f"\n{VERT}{GRAS}TOUS LES TESTS PASSENT{RESET}")
    return 0


# --- Main ---------------------------------------------------------------------

def main() -> int:
    parser = argparse.ArgumentParser(description="Tests exporter (multi-formes)")
    parser.add_argument('--template', type=str, default=None,
                        help='Chemin vers le template .xlsx')
    parser.add_argument('--avec-dim', type=str, default=None,
                        help='Fichier avec feuille de dimensions')
    parser.add_argument('--sans-dim', type=str, default=None,
                        help='Fichier sans feuille de dimensions')
    parser.add_argument('--forme', type=str, default='cylindre',
                        help='Forme à tester (défaut : cylindre)')
    args = parser.parse_args()

    # Template
    if args.template:
        template = Path(args.template)
        if not template.exists():
            print(f"{ROUGE}Template introuvable : {template}{RESET}")
            return 2
    else:
        template = _trouver_template()
        if template is None:
            print(f"{ROUGE}Aucun template trouvé.{RESET}")
            print("Cherchés dans :")
            for d in DOSSIER_TEMPLATES:
                print(f"  - {d}")
            print("Utilise --template pour préciser.")
            return 2

    # Fichiers de test
    avec, sans = _choisir_fichiers(args)

    # Dossier de sortie
    DOSSIER_SORTIE.mkdir(parents=True, exist_ok=True)

    print(f"{GRAS}Template                :{RESET} {template}")
    print(f"{GRAS}Fichier avec dimensions :{RESET} {avec or JAUNE + 'aucun' + RESET}")
    print(f"{GRAS}Fichier sans dimensions :{RESET} {sans or JAUNE + 'aucun' + RESET}")
    print(f"{GRAS}Forme testée            :{RESET} {args.forme}")
    print(f"{GRAS}Dossier de sortie       :{RESET} {DOSSIER_SORTIE}")

    if avec:
        test_01_avec_dimensions(template, avec, args.forme)
        test_03_mauvais_forme(template, avec, args.forme)
    else:
        _titre("Test 1 et 3 — Aucun fichier avec dimensions")
        _skip("tests 1 et 3", "aucun fichier fourni")

    if sans:
        test_02_sans_dimensions(template, sans)
    else:
        _titre("Test 2 — Aucun fichier sans dimensions")
        _skip("test 2", "aucun fichier fourni")

    return _resume()


if __name__ == '__main__':
    sys.exit(main())