"""
Tests de reader.py (multi-formes) — adapté au projet PycnoLab.

Structure attendue :
    pycnolab/
        pycno_core/moteur_python/import_excel/reader.py
        essais/bruts/*.xlsx
        test_reader.py

Usage :
    python test_reader.py
    python test_reader.py --fichier-avec-dim "essais/bruts/mon_fichier.xlsx"
    python test_reader.py --fichier-sans-dim "essais/bruts/autre.xlsx"
    python test_reader.py --forme rectangle
"""

from __future__ import annotations

import argparse
import sys
import traceback
from pathlib import Path

# --- Import du module reader ------------------------------------------------

RACINE = Path(__file__).parent
sys.path.insert(0, str(RACINE / 'pycno_core'))

try:
    from moteur_python.import_excel.reader import (
        read_workbook,
        lire_donnees_brutes,
        _lire_dimensions_pour_forme,
        _lire_diametres_generiques,
        _trouver_feuille_dimensions,
        _FORMES_VALIDES,
    )
except ImportError as e:
    print(f"[FATAL] Import impossible : {e}")
    print("Cherche reader.py :")
    for p in RACINE.rglob('reader.py'):
        print(f"  {p}")
    sys.exit(1)

from openpyxl import load_workbook


# --- Affichage ---------------------------------------------------------------

VERT = "\033[92m"
ROUGE = "\033[91m"
JAUNE = "\033[93m"
GRIS = "\033[90m"
GRAS = "\033[1m"
RESET = "\033[0m"

_resultats: list[tuple[str, str, str]] = []


def _titre(t: str):
    print()
    print(f"{GRAS}{'─' * 70}{RESET}")
    print(f"{GRAS}  {t}{RESET}")
    print(f"{GRAS}{'─' * 70}{RESET}")


def _ok(nom: str, msg: str = ""):
    _resultats.append((nom, "OK", msg))
    print(f"  {VERT}[OK]{RESET}    {nom}" + (f"  {GRIS}{msg}{RESET}" if msg else ""))


def _echec(nom: str, msg: str):
    _resultats.append((nom, "ECHEC", msg))
    print(f"  {ROUGE}[ECHEC]{RESET} {nom}")
    print(f"          {ROUGE}{msg}{RESET}")


def _skip(nom: str, raison: str):
    _resultats.append((nom, "SKIP", raison))
    print(f"  {JAUNE}[SKIP]{RESET}  {nom}  {GRIS}({raison}){RESET}")


def _info(msg: str):
    print(f"  {GRIS}{msg}{RESET}")


# --- Détection des fichiers --------------------------------------------------

DOSSIER_BRUTS = RACINE / 'essais' / 'bruts'


def _lister_fichiers() -> list[Path]:
    if not DOSSIER_BRUTS.exists():
        return []
    return sorted(DOSSIER_BRUTS.glob('*.xlsx'))


def _a_feuille_dimensions(chemin: Path) -> bool:
    try:
        wb = load_workbook(chemin, data_only=True, read_only=True)
        try:
            return _trouver_feuille_dimensions(wb) is not None
        finally:
            wb.close()
    except Exception:
        return False


def _choisir_fichiers(args) -> tuple[Path | None, Path | None]:
    if args.fichier_avec_dim:
        c = Path(args.fichier_avec_dim)
        if not c.exists():
            print(f"{ROUGE}Fichier introuvable : {c}{RESET}")
            sys.exit(2)
        avec = c
    else:
        avec = None

    if args.fichier_sans_dim:
        c = Path(args.fichier_sans_dim)
        if not c.exists():
            print(f"{ROUGE}Fichier introuvable : {c}{RESET}")
            sys.exit(2)
        sans = c
    else:
        sans = None

    if avec is None or sans is None:
        candidats = _lister_fichiers()
        if not candidats:
            print(f"{ROUGE}Aucun .xlsx dans {DOSSIER_BRUTS}{RESET}")
            sys.exit(2)
        _info(f"Analyse de {len(candidats)} fichier(s) dans {DOSSIER_BRUTS}…")
        for f in candidats:
            a_dim = _a_feuille_dimensions(f)
            tag = "avec dim" if a_dim else "sans dim"
            _info(f"  - {f.name}  [{tag}]")
            if avec is None and a_dim:
                avec = f
            if sans is None and not a_dim:
                sans = f

    return avec, sans


# --- Tests -------------------------------------------------------------------

def test_01_detection_feuille(chemin: Path):
    _titre("Test 1 — Détection de la feuille de dimensions")
    try:
        wb = load_workbook(chemin, data_only=True, read_only=True)
        try:
            nom = _trouver_feuille_dimensions(wb)
            if nom is None:
                _echec("_trouver_feuille_dimensions", "aucune feuille détectée")
                return
            _ok("_trouver_feuille_dimensions", f"feuille : {nom!r}")
        finally:
            wb.close()
    except Exception as e:
        _echec("_trouver_feuille_dimensions", f"{type(e).__name__} : {e}")


def test_02_donnees_brutes(chemin: Path):
    _titre("Test 2 — Lecture brute + flag a_feuille_dimensions")
    try:
        info = lire_donnees_brutes(str(chemin), "traction")
        flag = info.get("a_feuille_dimensions")
        gammes = info.get("gammes", [])
        if not isinstance(flag, bool):
            _echec("a_feuille_dimensions", f"attendu bool, reçu {type(flag).__name__}")
            return
        _ok("a_feuille_dimensions", f"= {flag}")
        if not gammes:
            _echec("gammes", "aucune gamme détectée")
            return
        _ok("gammes", f"{len(gammes)} : {[g['nom_gamme'] for g in gammes]}")
        for g in gammes:
            _info(f"  - {g['nom_gamme']} : {len(g['echantillons'])} éch.")
    except Exception as e:
        _echec("lire_donnees_brutes", f"{type(e).__name__} : {e}")
        traceback.print_exc()


def test_03_lecture_forme(chemin: Path, forme: str):
    _titre(f"Test 3 — Lecture de la forme {forme!r}")
    try:
        wb = load_workbook(chemin, data_only=True, read_only=True)
        try:
            dims = _lire_dimensions_pour_forme(wb, forme)
        finally:
            wb.close()
        if not dims:
            _skip(f"lecture {forme}", "aucun bloc trouvé")
            return
        _ok(f"lecture {forme}", f"{len(dims)} bloc(s)")
        for nom_gamme, ech in dims.items():
            _info(f"  - {nom_gamme} : {len(ech)} éch.")
            for e in ech[:2]:
                _info(
                    f"      {e['identifiant']}  forme={e['forme']}  "
                    f"S={e['section_mm2']:.4f} mm²"
                )
    except Exception as e:
        _echec(f"lecture {forme}", f"{type(e).__name__} : {e}")
        traceback.print_exc()


def test_04_read_workbook(chemin: Path, forme: str):
    _titre(f"Test 4 — read_workbook avec feuille (forme={forme!r})")
    try:
        res = read_workbook(str(chemin), "traction", forme=forme)
        total = 0
        for g in res["gammes"]:
            for e in g["echantillons"]:
                if "section_mm2" not in e or e["section_mm2"] <= 0:
                    _echec(
                        f"section de {e['identifiant']}",
                        f"manquante ou nulle : {e.get('section_mm2')}",
                    )
                    return
                total += 1
        _ok("read_workbook", f"{total} éch. avec section valide")
        for g in res["gammes"]:
            e = g["echantillons"][0]
            _info(
                f"  - {g['nom_gamme']} : {e['identifiant']}  "
                f"forme={e.get('forme')}  S={e['section_mm2']:.4f}"
            )
    except Exception as e:
        _echec("read_workbook", f"{type(e).__name__} : {e}")
        traceback.print_exc()


def test_05_mauvais_forme(chemin: Path, bonne: str):
    _titre("Test 5 — Mauvais choix de forme (doit lever)")
    autres = [f for f in _FORMES_VALIDES if f != bonne]
    if not autres:
        _skip("mauvais choix", "aucune autre forme")
        return
    mauvaise = autres[0]
    try:
        read_workbook(str(chemin), "traction", forme=mauvaise)
        _echec("erreur attendue", f"aucune erreur pour forme={mauvaise!r}")
    except ValueError as e:
        msg = str(e)
        if any(k in msg for k in ("correspondant", "introuvable", "Aucune")):
            _ok("erreur claire", f"{msg[:70]}…")
        else:
            _echec("message inattendu", msg)
    except Exception as e:
        _echec("mauvaise exception", f"{type(e).__name__} : {e}")


def test_06_sans_feuille_geometrie(chemin: Path):
    _titre("Test 6 — Fichier sans feuille + géométrie UI")
    try:
        res = read_workbook(
            str(chemin),
            "traction",
            geometrie={"section_mm2": 40.0, "l0_mm": 20.0},
            forme="rectangle",
        )
        total = 0
        for g in res["gammes"]:
            for e in g["echantillons"]:
                if abs(e["section_mm2"] - 40.0) > 1e-9:
                    _echec(
                        f"section de {e['identifiant']}",
                        f"attendu 40.0, reçu {e['section_mm2']}",
                    )
                    return
                total += 1
        _ok("sections injectées", f"{total} éch. à 40.0 mm²")
    except Exception as e:
        _echec("read_workbook sans feuille", f"{type(e).__name__} : {e}")
        traceback.print_exc()


def test_07_sans_feuille_rien(chemin: Path):
    _titre("Test 7 — Fichier sans feuille + pas de géométrie (doit lever)")
    try:
        read_workbook(str(chemin), "traction", forme="rectangle")
        _echec("erreur attendue", "aucune erreur levée")
    except ValueError as e:
        msg = str(e)
        if "Pas de feuille" in msg or "section saisie" in msg:
            _ok("erreur claire", f"{msg[:70]}…")
        else:
            _echec("message inattendu", msg)
    except Exception as e:
        _echec("mauvaise exception", f"{type(e).__name__} : {e}")


def test_08_retrocompat(chemin: Path):
    _titre("Test 8 — Rétrocompatibilité _lire_diametres_generiques")
    try:
        wb = load_workbook(chemin, data_only=True, read_only=True)
        try:
            dims = _lire_diametres_generiques(wb)
        finally:
            wb.close()
        if not dims:
            _skip("_lire_diametres_generiques", "aucun bloc cylindre")
            return
        _ok("_lire_diametres_generiques", f"{len(dims)} bloc(s)")
        for nom, ech in dims.items():
            for e in ech[:1]:
                if "diametres_mm" not in e:
                    _echec(
                        f"diametres_mm de {e['identifiant']}",
                        "clé manquante",
                    )
                    return
        _ok("clé diametres_mm", "compat export OK")
    except Exception as e:
        _echec("_lire_diametres_generiques", f"{type(e).__name__} : {e}")
        traceback.print_exc()


# --- Résumé ------------------------------------------------------------------

def _resume() -> int:
    _titre("RÉSUMÉ")
    total = len(_resultats)
    ok = sum(1 for _, s, _ in _resultats if s == "OK")
    ko = sum(1 for _, s, _ in _resultats if s == "ECHEC")
    sk = sum(1 for _, s, _ in _resultats if s == "SKIP")
    print(f"  Total : {total}")
    print(f"  {VERT}OK{RESET}    : {ok}")
    print(f"  {ROUGE}ECHEC{RESET} : {ko}")
    print(f"  {JAUNE}SKIP{RESET}  : {sk}")
    if ko:
        print()
        print(f"  {ROUGE}En échec :{RESET}")
        for nom, s, msg in _resultats:
            if s == "ECHEC":
                print(f"    - {nom}  {GRIS}({msg}){RESET}")
        print(f"\n{ROUGE}{GRAS}ÉCHEC GLOBAL{RESET}")
        return 1
    print(f"\n{VERT}{GRAS}TOUS LES TESTS PASSENT{RESET}")
    return 0


# --- Main --------------------------------------------------------------------

def main() -> int:
    parser = argparse.ArgumentParser(description="Tests reader.py (multi-formes)")
    parser.add_argument("--fichier-avec-dim", type=str, default=None)
    parser.add_argument("--fichier-sans-dim", type=str, default=None)
    parser.add_argument("--forme", type=str, default="cylindre")
    args = parser.parse_args()

    avec, sans = _choisir_fichiers(args)

    print(f"{GRAS}Fichier avec dimensions :{RESET} {avec or JAUNE + 'aucun' + RESET}")
    print(f"{GRAS}Fichier sans dimensions :{RESET} {sans or JAUNE + 'aucun' + RESET}")
    print(f"{GRAS}Forme testée           :{RESET} {args.forme}")

    if avec:
        test_01_detection_feuille(avec)
        test_02_donnees_brutes(avec)
        test_03_lecture_forme(avec, args.forme)
        test_04_read_workbook(avec, args.forme)
        test_05_mauvais_forme(avec, args.forme)
        test_08_retrocompat(avec)
    else:
        _titre("Tests 1–5, 8 — Aucun fichier avec dimensions")
        _skip("tests avec dimensions", "aucun fichier fourni")

    if sans:
        test_06_sans_feuille_geometrie(sans)
        test_07_sans_feuille_rien(sans)
    else:
        _titre("Tests 6–7 — Aucun fichier sans dimensions")
        _skip("tests sans dimensions", "aucun fichier fourni")

    return _resume()


if __name__ == "__main__":
    sys.exit(main())