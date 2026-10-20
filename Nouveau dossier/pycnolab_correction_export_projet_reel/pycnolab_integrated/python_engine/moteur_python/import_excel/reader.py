from __future__ import annotations
from pathlib import Path
from typing import Any
from openpyxl import load_workbook
import math, re

DIAMETER_SHEET = 'DIAMETRE (CR,CS)'
IGNORED_SHEETS = {'INFO', 'RECCAPITULATIF'}


def _num(v):
    if v is None: return None
    if isinstance(v, (int, float)): return float(v)
    try: return float(str(v).strip().replace(',', '.'))
    except (TypeError, ValueError): return None


def _norm(v):
    return re.sub(r'[^A-Z0-9]+', '', str(v or '').upper())


def _sample_id(header, fallback):
    m = re.search(r'(?:ECH|ECHANTILLON|SAMPLE)\s*[-_ ]*([0-9]+)', str(header or ''), re.I)
    return f'ECH{int(m.group(1))}' if m else f'ECH{fallback}'


def _diameter_groups(ws):
    groups=[]
    # Les groupes sont identifiés par la présence de D_1, D_2, D_3 sur la ligne 2.
    for c in range(1, ws.max_column-2):
        labels=' '.join(str(ws.cell(2,c+i).value or '').upper() for i in range(3))
        if not ('D_1' in labels or 'D1' in labels) or not ('D_2' in labels or 'D2' in labels) or not ('D_3' in labels or 'D3' in labels):
            continue
        title=ws.cell(1,c).value
        if title is None:
            for p in range(c-1,0,-1):
                if ws.cell(1,p).value is not None:
                    title=ws.cell(1,p).value; break
        name=str(title or f'GAMME_{len(groups)+1}').strip()
        groups.append((name,c))
    unique=[]
    seen=set()
    for name,c in groups:
        key=(name.upper(),c)
        if key not in seen: unique.append((name,c)); seen.add(key)
    return unique


def _horizontal_blocks(ws):
    blocks=[]
    for c in range(1, ws.max_column+1):
        h=str(ws.cell(1,c).value or '')
        if not re.search(r'ELONG',h,re.I): continue
        force=None
        for j in range(c+1,min(ws.max_column,c+8)+1):
            if re.search(r'FORCE',str(ws.cell(1,j).value or ''),re.I): force=j; break
        if force:
            blocks.append((_sample_id(h,len(blocks)+1),c,force))
    unique={sid:(sid,ec,fc) for sid,ec,fc in blocks}
    return list(unique.values())


def _vertical_data(ws):
    headers=[str(ws.cell(1,c).value or '') for c in range(1,ws.max_column+1)]
    def col(patterns):
        for c,h in enumerate(headers,1):
            if any(re.search(p,h,re.I) for p in patterns): return c
        return None
    idc=col([r'ID\s*ECHANT',r'IDENT'])
    elongc=col([r'ELONG',r'DEPLAC'])
    forcec=col([r'^FORCE',r'FORCE\s*\('])
    if not (idc and elongc and forcec): return {}
    curves={}
    for r in range(2,ws.max_row+1):
        sid=ws.cell(r,idc).value; elong=_num(ws.cell(r,elongc).value); force=_num(ws.cell(r,forcec).value)
        if sid is None or elong is None or force is None: continue
        curves.setdefault(str(sid).strip().upper(),[]).append({'deplacement_mm':elong,'force_newton':force})
    return curves


def read_workbook(path: str | Path) -> dict[str, Any]:
    path=Path(path)
    if not path.is_file(): raise FileNotFoundError(path)
    wb=load_workbook(path,data_only=True,read_only=False)
    if DIAMETER_SHEET not in wb.sheetnames: raise ValueError(f'Feuille obligatoire absente: {DIAMETER_SHEET}')
    dws=wb[DIAMETER_SHEET]
    dgroups=_diameter_groups(dws)
    if not dgroups: raise ValueError('Aucun groupe de diamètres détecté dans DIAMETRE (CR,CS).')
    diameter_data={}
    diameter_order=[]
    for name,start in dgroups:
        rows=[]
        for r in range(3,dws.max_row+1):
            ident=dws.cell(r,start-1).value
            vals=[_num(dws.cell(r,start+i).value) for i in range(3)]
            if ident is None: continue
            if any(v is None for v in vals): raise ValueError(f'{name}: diamètres incomplets pour {ident}')
            dm=sum(vals)/3; section=math.pi*(dm/2)**2
            rows.append({'identifiant':str(ident).strip().upper(),'diametres_mm':vals,'diametre_moyen_mm':dm,'section_mm2':section})
        key=_norm(name)
        diameter_data[key]=rows
        diameter_order.append((name, rows))
    gammes=[]
    for sheet_name in wb.sheetnames:
        if sheet_name in IGNORED_SHEETS or sheet_name==DIAMETER_SHEET: continue
        ws=wb[sheet_name]; blocks=_horizontal_blocks(ws)
        curves={}
        for sid,ec,fc in blocks:
            pts=[]
            for r in range(2,ws.max_row+1):
                x=_num(ws.cell(r,ec).value); y=_num(ws.cell(r,fc).value)
                if x is None and y is None: continue
                if x is not None and y is not None: pts.append({'deplacement_mm':x,'force_newton':y})
            if pts: curves[sid]=pts
        if not curves:
            curves=_vertical_data(ws)
        if not curves: continue
        drows=diameter_data.get(_norm(sheet_name),[])
        if not drows:
            sn=_norm(sheet_name)
            for title, rows in diameter_order:
                tn=_norm(title)
                if sn in tn or tn in sn:
                    drows=rows; break
        if not drows:
            # Dernier secours : même ordre de découverte entre feuilles de gammes et groupes de diamètres.
            discovered=[n for n in wb.sheetnames if n not in IGNORED_SHEETS and n != DIAMETER_SHEET]
            try: drows=diameter_order[discovered.index(sheet_name)][1]
            except (ValueError, IndexError): drows=[]
        dmap={x['identifiant']:x for x in drows}
        samples=[]
        for sid,pts in curves.items():
            d=dmap.get(sid)
            if d is None: continue
            samples.append({**d,'points_courbe':pts})
        if samples:
            gammes.append({'nom_gamme':sheet_name,'echantillons':samples,'diametre_moyen_global_mm':sum(s['diametre_moyen_mm'] for s in samples)/len(samples)})
    if not gammes: raise ValueError('Aucune gamme avec courbe exploitable détectée.')
    return {'gammes':gammes,'source':str(path)}


def lire_classeur_complet(path: str | Path) -> dict[str, Any]:
    return read_workbook(path)


def lire_gamme_complete(chemin_fichier: str | Path) -> dict[str, Any]:
    """Adaptateur de compatibilité pour les routes historiques mono-gamme.

    La nouvelle lecture utilise read_workbook() et découvre toutes les gammes.
    Cette fonction retourne la première gamme dans l'ancien format attendu par
    _from_excel et les anciennes routes API.
    """
    data = read_workbook(chemin_fichier)
    if not data['gammes']:
        raise ValueError('Aucune gamme exploitable dans le classeur.')
    gamme = data['gammes'][0]
    return {
        'materiau_nom': gamme['nom_gamme'],
        'norme_code': 'À définir',
        'echantillons': [
            {
                **sample,
                'longueur_initiale_mm': 20.0,
                'force_rupture_newton': max((p['force_newton'] for p in sample['points_courbe']), default=None),
                'deplacement_rupture_mm': next((p['deplacement_mm'] for p in sample['points_courbe'] if p['force_newton'] == max(q['force_newton'] for q in sample['points_courbe'])), None) if sample['points_courbe'] else None,
            }
            for sample in gamme['echantillons']
        ],
        'calculs_demandes': ['energie_rupture', 'limite_elastique'],
    }
