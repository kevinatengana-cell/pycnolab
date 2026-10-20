from openpyxl import load_workbook
from copy import copy
from pathlib import Path
from production_workbook_exporter import process_workbook
from classeur_entree_reader import read_workbook
src='/home/ubuntu/TRACTIONFIBREVEGETALE_entree_3_feuilles.xlsx'
case='/home/ubuntu/entree_test_3_gammes.xlsx'
wb=load_workbook(src)
ws=wb['COROUI']; extra=wb.copy_worksheet(ws); extra.title='GAMME_3'
# Ne garder que 5 blocs dans cette gamme dérivée.
for c in range(41, extra.max_column+1):
    for r in range(1, extra.max_row+1): extra.cell(r,c).value=None
# Ajouter un troisième groupe de diamètres à partir de COROUI.
d=wb['DIAMETRE (CR,CS)']; start=18
d.cell(1,start).value='DIAMETRE GAMME_3'
for off,label in enumerate(['D_1(mm)','D_2(mm)','D_3(mm)','D_Moy(mm)','S_(mm2)','D_Moy Global(mm)']): d.cell(2,start+off).value=label
for r in range(3,8):
    d.cell(r,start-1).value=d.cell(r,1).value
    for off in range(3): d.cell(r,start+off).value=d.cell(r,2+off).value
wb.save(case)
data=read_workbook(case)
assert [g['nom_gamme'] for g in data['gammes']]==['COROUI','COSOU','GAMME_3']
assert [len(g['echantillons']) for g in data['gammes']]==[22,22,5]
out='/home/ubuntu/RAPPORT_TRACTION_3_GAMMES.xlsx'
process_workbook(case,'/home/ubuntu/upload/TRACTIONFIBREVEGETALE.xlsx',out)
check=load_workbook(out,data_only=False)
assert 'GAMME_3' in check.sheetnames
assert len(check['GAMME_3']._charts)==5
print('OK gammes=',[g['nom_gamme'] for g in data['gammes']], 'echantillons=',[len(g['echantillons']) for g in data['gammes']])
